import contextlib
import copy
import csv
import io
import json
import os
import shutil
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from unittest.mock import patch
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, urlsplit

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skills/timesheet-copilot"
sys.path.insert(0, str(SKILL / "scripts"))

from config import TimesheetError, load_profile, period, select_projects, week_segments
from pms import NoRedirect, PMS, credential
from report import HEADER, validate
from timesheet import main, prepare
from openpyxl import load_workbook

FIRST, LAST = date(2026, 9, 13), date(2026, 9, 19)
CSV = (SKILL / "examples/pms-export.csv").read_text()
CONTROL = [{"start": str(FIRST), "end": str(LAST), "minutes": 275}]


def as_csv(rows):
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows(rows)
    return stream.getvalue()


def mutate(row, column, value):
    rows = list(csv.reader(io.StringIO(CSV)))
    rows[row][column] = value
    return as_csv(rows)


class TempCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name).resolve()
        self.profile_path = self.base / "profile.json"
        self.profile_path.write_text((SKILL / "examples/profile.example.json").read_text())
        self.profile = load_profile(self.profile_path)
        self.project = self.profile["projects"][0]


class ConfigTests(TempCase):
    def test_relative_output_uses_profile_location(self):
        self.assertEqual(Path(self.profile["output_dir"]), self.base / "timesheets")

    def test_previous_complete_week_at_boundary(self):
        for today in (date(2026, 9, 20), date(2026, 9, 21), date(2026, 9, 26)):
            self.assertEqual(period(self.profile, today=today), (FIRST, LAST))
        self.profile["week_starts_on"] = "monday"
        self.assertEqual(period(self.profile, today=date(2026, 9, 21)), (date(2026, 9, 14), date(2026, 9, 20)))

    def test_year_boundary(self):
        self.assertEqual(period(self.profile, today=date(2026, 1, 1)), (date(2025, 12, 21), date(2025, 12, 27)))

    def test_explicit_range_and_invalid_dates(self):
        self.assertEqual(period(self.profile, "2026-09-01", "2026-09-15"), (date(2026, 9, 1), date(2026, 9, 15)))
        for first, last in (("2026-09-01", None), ("bad", "bad"), ("2026-09-20", "2026-09-13")):
            with self.subTest(first=first, last=last), self.assertRaises(TimesheetError):
                period(self.profile, first, last)

    def test_scope_is_explicit(self):
        self.profile["projects"].append({"key": "archive", "id": 102, "name": "Archive", "active": False})
        self.assertEqual([p["key"] for p in select_projects(self.profile)], ["atlas"])
        self.assertEqual(select_projects(self.profile, ["archive"])[0]["id"], 102)
        with self.assertRaises(TimesheetError):
            select_projects(self.profile, ["unknown"])

    def test_invalid_configuration(self):
        original = json.loads(self.profile_path.read_text())
        mutations = [
            lambda p: p.update(source="jira"),
            lambda p: p.update(timezone="Invalid/Zone"),
            lambda p: p.update(week_starts_on="Sunday"),
            lambda p: p.update(week_starts_on=[]),
            lambda p: p.update(unknown=True),
            lambda p: p["pms"].update(base_url="http://pms.example.com"),
            lambda p: p["pms"].update(base_url="https://user:pass@pms.example.com"),
            lambda p: p["pms"].update(base_url="https://pms.example.com/api"),
            lambda p: p["pms"].update(base_url="https://pms.example.com:notaport"),
            lambda p: p["pms"].update(token="secret"),
            lambda p: p["pms"].update(token_file="token.txt"),
            lambda p: p["projects"].append(p["projects"][0].copy()),
            lambda p: p["projects"][0].update(id=True),
            lambda p: p["projects"][0].update(active="false"),
        ]
        for mutation in mutations:
            p = copy.deepcopy(original)
            mutation(p)
            self.profile_path.write_text(json.dumps(p))
            with self.subTest(profile=p), self.assertRaises(TimesheetError):
                load_profile(self.profile_path)


class ReportTests(unittest.TestCase):
    def run_report(self, text=CSV, controls=CONTROL, first=FIRST, last=LAST):
        return validate(text, "Example Atlas", first, last, controls)

    def test_detail_sum_excludes_subtotals(self):
        report = self.run_report()
        self.assertEqual(report.total_minutes, 275)
        self.assertEqual(len(report.entries), 3)
        self.assertEqual(report.subtotal_rows, [4, 6, 7])
        self.assertEqual(report.warnings, [])

    def test_multiline_and_formula_like_text_are_preserved(self):
        report = self.run_report(mutate(1, 3, '=HYPERLINK("https://example.com")\nSecond line, exact text'))
        self.assertEqual(report.rows[1][3], '=HYPERLINK("https://example.com")\nSecond line, exact text')

    def test_mismatched_daily_weekly_and_api_totals(self):
        for text in (mutate(3, 4, "9"), mutate(6, 5, "36")):
            with self.subTest(text=text), self.assertRaisesRegex(TimesheetError, "total does not match"):
                self.run_report(text)
        bad = [{**CONTROL[0], "minutes": 276}]
        with self.assertRaisesRegex(TimesheetError, "PMS total mismatch"):
            self.run_report(controls=bad)

    def test_scope_and_invalid_rows(self):
        cases = [(1, 1, "Other Project"), (1, 0, "2026-09-12"), (1, 0, "not a date"),
                 (1, 4, "1.5"), (1, 4, "-1"), (1, 5, "60"), (1, 6, ""),
                 (1, 0, ""), (1, 3, "bad\x01control")]
        for row, col, value in cases:
            with self.subTest(row=row, col=col, value=value), self.assertRaises(TimesheetError):
                self.run_report(mutate(row, col, value))

    def test_empty_or_login_response_and_extra_columns_fail(self):
        for text in ("", "<html>Please sign in</html>", CSV.replace("Resource", "Resource,Secret"),
                     CSV.replace("Alex Example", "Alex Example,Extra")):
            with self.subTest(text=text), self.assertRaises(TimesheetError):
                self.run_report(text)

    def test_empty_requires_independent_zero(self):
        header = as_csv([HEADER])
        with self.assertRaises(TimesheetError):
            self.run_report(header, controls=None)
        report = self.run_report(header, controls=[{**CONTROL[0], "minutes": 0}])
        self.assertEqual(report.total_minutes, 0)
        with self.assertRaises(TimesheetError):
            self.run_report(header)

    def test_missing_control_is_visible(self):
        report = self.run_report(controls=None)
        self.assertIn("No independent PMS total", " ".join(report.warnings))

    def test_resource_description_total_is_still_an_entry(self):
        report = self.run_report(mutate(1, 3, "Total"))
        self.assertEqual(len(report.entries), 3)

    def test_multweek_controls_and_date_coverage(self):
        rows = list(csv.reader(io.StringIO(CSV)))
        rows.extend([["2026-09-20", "Example Atlas", "Testing", "Another week", "1", "0", "Alex Example"],
                     ["", "", "", "Total", "1", "0", ""], ["", "", "", "Weekly Total", "1", "0", ""]])
        controls = CONTROL + [{"start": "2026-09-20", "end": "2026-09-26", "minutes": 60}]
        report = self.run_report(as_csv(rows), controls, last=date(2026, 9, 26))
        self.assertEqual(report.total_minutes, 335)
        for bad in ([], CONTROL + CONTROL, [{**CONTROL[0], "end": "2026-09-18"}]):
            with self.subTest(bad=bad), self.assertRaises(TimesheetError):
                self.run_report(controls=bad)

    def test_repeated_and_misdated_subtotals_fail(self):
        rows = list(csv.reader(io.StringIO(CSV)))
        rows.append(next(row for row in rows if len(row) == 7 and row[3] == "Weekly Total"))
        with self.assertRaises(TimesheetError):
            self.run_report(as_csv(rows))
        with self.assertRaises(TimesheetError):
            self.run_report(mutate(3, 0, "2026-09-14"))

    def test_other_date_formats_and_bom(self):
        for day in ("09/13/2026", "2026-09-13", "Sunday, 13 September, 2026", "Sunday, 13 September 2026"):
            self.assertEqual(self.run_report("\ufeff" + mutate(1, 0, day)).total_minutes, 275)

    def test_empty_duplicate_totals_fail(self):
        rows = [HEADER, ["", "", "", "Weekly Total", "0", "0", ""], ["", "", "", "Weekly Total", "0", "0", ""]]
        with self.assertRaises(TimesheetError):
            self.run_report(as_csv(rows), controls=[{**CONTROL[0], "minutes": 0}])

    def test_missing_last_daily_total_is_visible(self):
        rows = list(csv.reader(io.StringIO(CSV)))
        del rows[5]
        self.assertIn("Some days have no CSV subtotal", " ".join(self.run_report(as_csv(rows)).warnings))


class WorkbookTests(TempCase):
    def test_saved_workbook_preserves_values_and_layout(self):
        text = mutate(1, 3, '=Literal source text\nA longer second line with commas, exactly preserved.')
        result = prepare(self.profile, self.project, FIRST, LAST, text, CONTROL)
        path = Path(result["workbook"])
        self.assertTrue(path.exists())
        self.assertEqual(result["hours"], "4h 35m")
        self.assertEqual((path.parent / "source.csv").read_bytes(), text.encode("utf-8"))
        self.assertEqual(json.loads((path.parent / "verification.json").read_text())["total_minutes"], 275)
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        ws = wb.active
        self.assertEqual(ws["D2"].value, '=Literal source text\nA longer second line with commas, exactly preserved.')
        self.assertEqual(ws["D2"].data_type, "s")
        self.assertEqual(ws["E2"].value, 2)
        self.assertEqual(ws["F2"].value, 30)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws["D4"].font.bold)
        self.assertGreater(ws.row_dimensions[2].height, 25)
        for row in ws:
            for cell in row:
                self.assertTrue(cell.alignment.wrap_text)

    def test_rerun_preserves_previous_workbook(self):
        a = prepare(self.profile, self.project, FIRST, LAST, CSV, CONTROL)
        original = Path(a["workbook"]).read_bytes()
        b = prepare(self.profile, self.project, FIRST, LAST, CSV, CONTROL)
        self.assertNotEqual(a["workbook"], b["workbook"])
        self.assertEqual(original, Path(a["workbook"]).read_bytes())

    def test_failure_does_not_leave_final_files(self):
        for text in (mutate(1, 3, "word " * 1000), mutate(3, 4, "9")):
            with self.subTest(text_length=len(text)), self.assertRaises(TimesheetError):
                prepare(self.profile, self.project, FIRST, LAST, text, CONTROL)
        output = Path(self.profile["output_dir"])
        self.assertFalse(list(output.rglob("*.xlsx")))

    def test_filename_cannot_escape_output_directory(self):
        self.profile["filename_prefix"] = "../../Team"
        self.project["label"] = "../Project:Name"
        result = prepare(self.profile, self.project, FIRST, LAST, CSV, CONTROL)
        self.assertTrue(Path(result["workbook"]).is_relative_to(Path(self.profile["output_dir"])))


class PMSTests(unittest.TestCase):
    def client(self):
        with patch.dict(os.environ, {"TEST_PMS_TOKEN": "secret-for-tests"}):
            return PMS({"base_url": "https://pms.example.com", "token_env": "TEST_PMS_TOKEN"})

    def test_fetch_checks_every_week_segment(self):
        client = self.client()
        with patch.object(client, "get", side_effect=[CSV, '{"weeklyTotalInMinute": 275}', '{"weeklyTotalInMinute": 0}']) as get:
            text, controls = client.fetch(101, FIRST, date(2026, 9, 22), "sunday")
        self.assertEqual(text, CSV)
        self.assertEqual([c["minutes"] for c in controls], [275, 0])
        self.assertEqual(controls[-1]["end"], "2026-09-22")
        self.assertEqual(get.call_args_list[0].args[1]["projectIds"], "101")
        self.assertEqual(get.call_args_list[1].args[1]["endDate"], "2026-09-19")
        self.assertEqual(get.call_args_list[2].args[1]["startDate"], "2026-09-20")

    def test_missing_and_invalid_totals_are_not_zero(self):
        for payload in ("{}", '{"weeklyTotalInMinute": null}', '{"weeklyTotalInMinute": "275"}',
                        '{"weeklyTotalInMinute": true}', '{"weeklyTotalInMinute": -1}', "[]", "login HTML"):
            client = self.client()
            with self.subTest(payload=payload), patch.object(client, "get", side_effect=[CSV, payload]), self.assertRaises(TimesheetError):
                client.fetch(101, FIRST, LAST, "sunday")

    def test_auth_failure_has_no_secret_or_response_body(self):
        client = self.client()
        for code in (401, 403, 302, 429, 500):
            error = HTTPError("https://example.com/?token=secret-for-tests", code, "secret-for-tests", {}, None)
            with self.subTest(code=code), patch.object(client._opener, "open", side_effect=error):
                with self.assertRaises(TimesheetError) as caught:
                    client.get("/api/timesheets", {})
                self.assertNotIn("secret-for-tests", str(caught.exception))

    def test_network_failure_and_redirect_refusal(self):
        client = self.client()
        with patch.object(client._opener, "open", side_effect=URLError("sensitive details")), self.assertRaises(TimesheetError):
            client.get("/api/timesheets", {})
        self.assertIsNone(NoRedirect().redirect_request(None, None, 302, "", {}, "https://elsewhere.example"))

    def test_real_request_shape_is_get_and_bounded(self):
        client = self.client()
        response = io.BytesIO(b'{"weeklyTotalInMinute":0}')
        with patch.object(client._opener, "open", return_value=response) as opened:
            self.assertIn("weeklyTotalInMinute", client.get("/api/timesheets", {"projectId": "101"}))
        request = opened.call_args.args[0]
        self.assertEqual(request.get_method(), "GET")
        self.assertEqual(parse_qs(urlsplit(request.full_url).query), {"projectId": ["101"]})
        self.assertEqual(request.get_header("Authorization"), "Bearer secret-for-tests")
        self.assertEqual(opened.call_args.kwargs["timeout"], 45)

    def test_partial_week_segments(self):
        self.assertEqual(list(week_segments(date(2026, 9, 18), date(2026, 9, 22), "sunday")),
                         [(date(2026, 9, 18), LAST), (date(2026, 9, 20), date(2026, 9, 22))])

    def test_credential_missing_and_file_permissions(self):
        with patch.dict(os.environ, {}, clear=True), self.assertRaises(TimesheetError):
            credential({"token_env": "PMS_TOKEN"})
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "credential"
            path.write_text("test-token")
            path.chmod(0o600)
            self.assertEqual(credential({"token_file": str(path)}), "test-token")
            if os.name == "posix":
                path.chmod(0o644)
                with self.assertRaises(TimesheetError):
                    credential({"token_file": str(path)})


class CLITests(TempCase):
    def test_copied_skill_runs_from_unrelated_folder(self):
        copied = self.base / "installed-skill"
        shutil.copytree(SKILL, copied, ignore=shutil.ignore_patterns("__pycache__"))
        run = subprocess.run([sys.executable, str(copied / "scripts/timesheet.py"), "format",
                              "--profile", str(self.profile_path), "--project", "atlas",
                              "--csv", str(copied / "examples/pms-export.csv"),
                              "--start", str(FIRST), "--end", str(LAST)],
                             cwd=self.base, capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        result = json.loads(run.stdout)["results"][0]
        self.assertIn("PMS total not checked", result["verification"])

    def test_demo_through_actual_entrypoint(self):
        command = [sys.executable, str(SKILL / "scripts/timesheet.py"), "format", "--profile", str(self.profile_path),
                   "--project", "atlas", "--csv", str(SKILL / "examples/pms-export.csv"),
                   "--start", str(FIRST), "--end", str(LAST), "--expected-total-minutes", "275"]
        run = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        payload = json.loads(run.stdout)
        self.assertEqual(payload["errors"], [])
        self.assertEqual(payload["results"][0]["total_minutes"], 275)
        self.assertTrue(Path(payload["results"][0]["workbook"]).exists())

    def test_init_does_not_overwrite(self):
        original = self.profile_path.read_bytes()
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(main(["init", "--profile", str(self.profile_path)]), 1)
        self.assertEqual(self.profile_path.read_bytes(), original)

    def test_partial_batch_reports_success_and_failure(self):
        raw = json.loads(self.profile_path.read_text())
        raw["pms"]["base_url"] = "https://pms.internal.example"
        raw["projects"].append({"key": "beta", "id": 102, "name": "Example Beta"})
        self.profile_path.write_text(json.dumps(raw))
        out = io.StringIO()
        with patch("timesheet.PMS") as provider, contextlib.redirect_stdout(out):
            provider.return_value.fetch.side_effect = [(CSV, CONTROL), TimesheetError("Access unavailable")]
            status = main(["fetch", "--profile", str(self.profile_path), "--start", str(FIRST), "--end", str(LAST)])
        result = json.loads(out.getvalue())
        self.assertEqual(status, 1)
        self.assertEqual(len(result["results"]), 1)
        self.assertEqual(result["errors"], [{"project": "beta", "error": "Access unavailable"}])


if __name__ == "__main__":
    unittest.main()
