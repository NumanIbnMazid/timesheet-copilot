"""One source-preserving Excel formatter, adapted from the original skill."""
from __future__ import annotations

import math
import textwrap
import unicodedata

from config import require


def visual_length(value):
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in str(value))


def write_workbook(report, path, first):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        from config import TimesheetError
        raise TimesheetError("The Excel helper is missing. Install the skill's requirements.txt in the assistant's Python environment.") from None
    wb = Workbook()
    ws = wb.active
    ws.title = "timesheet_" + first.strftime("%Y%m%d")
    mins = [18, 24, 18, 52, 9, 10, 20]
    caps = [36, 42, 30, 85, 12, 12, 34]
    widths = []
    for col in range(7):
        longest = max((max((visual_length(line) for line in str(row[col]).splitlines()), default=0)
                       for row in report.rows), default=0)
        width = min(caps[col], max(mins[col], longest * 1.12 + 3))
        widths.append(width)
        ws.column_dimensions[get_column_letter(col + 1)].width = width
    thin = Side(style="thin", color="D7DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row_num, row in enumerate(report.rows, 1):
        special = row_num == 1 or row_num in report.subtotal_rows
        height = 28 if row_num == 1 else 21
        for col, value in enumerate(row, 1):
            cell = ws.cell(row_num, col, value if value != "" else None)
            if isinstance(value, str) and value:
                # Exported task text beginning '=' must remain literal text.
                cell.data_type = "s"
            cell.font = Font(name="Aptos Narrow", size=12, bold=special,
                             color="FFFFFF" if row_num == 1 else "172B4D")
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if col in (5, 6) else "left")
            cell.border = border
            if row_num == 1:
                cell.fill = PatternFill("solid", fgColor="24476B")
            elif special:
                cell.fill = PatternFill("solid", fgColor="E8EDF3")
            elif row_num % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F8FA")
            if col in (5, 6) and row_num > 1:
                cell.number_format = "0"
            capacity = max(1, int((widths[col - 1] - 3) * 0.85))
            lines = sum(max(1, len(textwrap.wrap(line, capacity)), math.ceil(visual_length(line) / capacity))
                        for line in str(value).split("\n"))
            height = max(height, lines * 17 + 6)
        require(height <= 409, f"CSV row {row_num}: text is too long to display fully in an Excel row. Use a shorter source description or an agreed alternate layout.")
        ws.row_dimensions[row_num].height = height
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:G{ws.max_row}"
    wb.save(path)
    # Read back the actual saved bytes, not just the in-memory cells.
    saved = load_workbook(path)
    try:
        actual = saved.active
        for r, row in enumerate(report.rows, 1):
            for c, value in enumerate(row, 1):
                cell = actual.cell(r, c)
                expected = None if value == "" else value
                require(cell.value == expected and cell.data_type != "f", "Saved workbook differs from the source export.")
    finally:
        saved.close()
